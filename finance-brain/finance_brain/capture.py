"""Capture a workbook: tables, schema, relationships, formula lineage, fragility findings, notes.

Pure Python (openpyxl). No HydraDB, no LLM. Output is a JSON-serialisable dict (see capture()).
"""
from __future__ import annotations

import datetime as dt
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter

TOTAL_RE = re.compile(r"^\s*(grand\s+)?(sub-?)?total\b", re.I)
DERIVED_RE = re.compile(r"^\s*(gross\s+(profit|margin)|ebitda|operating\s+(income|profit|margin)|net\s+(income|profit|margin)|contribution\s+margin)\b", re.I)
PASTED_RE = re.compile(r"\b(pasted|copied|typed|not linked|hard-?coded|manual(ly)?( input)?)\b", re.I)
CELL_RE = re.compile(r"^([A-Z]{1,3})(\d+)$")
RANGE_RE = re.compile(r"^([A-Z]{1,3})(\d+):([A-Z]{1,3})(\d+)$")
COLS_RE = re.compile(r"^([A-Z]{1,3}):([A-Z]{1,3})$")
ROWS_RE = re.compile(r"^(\d+):(\d+)$")
NAME_RE = re.compile(r"^[A-Za-z_\\][A-Za-z0-9_.]*$")
CODE_RE = re.compile(r"^(?:[A-Z]{1,5}-?\d{2,}|\d{3,6}|[A-Z]{2,5}|L\d)$")
LOOKUP_FUNCS = {"VLOOKUP", "XLOOKUP", "HLOOKUP", "INDEX", "MATCH", "LOOKUP"}
BENIGN_CONSTANTS = {0, 1, 2, 12, 100, 1000, 10000, 1000000, 365, 360, 52, 4, 3}


def slug(s: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(s).lower()).strip("_") or "x"


def iso(v: Any) -> Any:
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()[:10]
    if isinstance(v, float) and v != v:  # NaN
        return None
    return v


@dataclass
class CellInfo:
    row: int
    col: int
    value: Any            # cached value (data_only load)
    formula: str | None
    fmt: str
    bold: bool
    filled: bool
    comment: str | None
    author: str | None

    @property
    def addr(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"

    @property
    def is_str(self) -> bool:
        return isinstance(self.value, str) and self.formula is None

    @property
    def is_num(self) -> bool:
        return isinstance(self.value, (int, float)) and not isinstance(self.value, bool)

    @property
    def is_date(self) -> bool:
        return isinstance(self.value, (dt.datetime, dt.date))


@dataclass
class Table:
    id: str
    slug: str
    sheet: str
    name: str
    title: str | None
    kind: str                 # table | kv | notes
    c1: int
    c2: int
    top_row: int
    header_row: int | None
    first_data_row: int | None
    last_row: int
    helper_rows: list[int] = field(default_factory=list)
    title_rows: list[int] = field(default_factory=list)
    derived_rows: dict[int, str] = field(default_factory=dict)
    columns: list[dict] = field(default_factory=list)
    rows: list[list] = field(default_factory=list)
    pairs: list[dict] = field(default_factory=list)
    text: list[str] = field(default_factory=list)
    formula_share: float = 0.0
    narrative: bool = False
    role: str = "data"

    @property
    def range(self) -> str:
        return f"{get_column_letter(self.c1)}{self.top_row}:{get_column_letter(self.c2)}{self.last_row}"

    def data_rows(self) -> list[int]:
        if self.first_data_row is None:
            return []
        return [r for r in range(self.first_data_row, self.last_row + 1) if r not in self.derived_rows]


# ----------------------------------------------------------------------------- loading

def load_grid(path: str | Path):
    wbf = openpyxl.load_workbook(path, data_only=False)
    wbv = openpyxl.load_workbook(path, data_only=True)
    grids: dict[str, dict[tuple[int, int], CellInfo]] = {}
    sheets_meta = []
    names: dict[str, str] = {}
    dn = wbf.defined_names
    try:
        items = list(dn.items())
    except AttributeError:  # openpyxl < 3.1
        items = [(d.name, d) for d in dn.definedName]
    for name, d in items:
        if getattr(d, "attr_text", None):
            names[name] = d.attr_text
    for wsf in wbf.worksheets:
        wsv = wbv[wsf.title]
        grid: dict[tuple[int, int], CellInfo] = {}
        n_formulas = 0
        try:
            sheet_names = list(wsf.defined_names.items())
        except Exception:
            sheet_names = []
        for name, d in sheet_names:
            if getattr(d, "attr_text", None):
                names[name] = d.attr_text
        for row in wsf.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                formula = c.value if isinstance(c.value, str) and c.value.startswith("=") else None
                if formula:
                    n_formulas += 1
                cached = wsv.cell(row=c.row, column=c.column).value if formula else c.value
                fill = c.fill
                filled = bool(fill is not None and fill.fill_type == "solid" and fill.fgColor is not None
                              and str(getattr(fill.fgColor, "rgb", "")) not in ("00000000", "FFFFFFFF", "None"))
                grid[(c.row, c.column)] = CellInfo(
                    c.row, c.column, cached, formula, c.number_format or "General", bool(c.font and c.font.bold), filled,
                    c.comment.text if c.comment else None, c.comment.author if c.comment else None)
        grids[wsf.title] = grid
        sheets_meta.append({
            "name": wsf.title, "state": wsf.sheet_state, "dims": wsf.dimensions, "n_cells": len(grid), "n_formulas": n_formulas,
            "hidden_rows": [r for r, d in wsf.row_dimensions.items() if d.hidden],
            "hidden_cols": [k for k, d in wsf.column_dimensions.items() if d.hidden],
            "merged": [str(m) for m in wsf.merged_cells.ranges],
        })
    return grids, sheets_meta, names


# ----------------------------------------------------------------------------- block and table detection

def find_blocks(grid: dict[tuple[int, int], CellInfo]) -> list[tuple[int, int, int, int]]:
    """Rectangular blocks separated by fully blank rows, then fully blank columns."""
    if not grid:
        return []
    rows = sorted({r for r, _ in grid})
    runs, start, prev = [], rows[0], rows[0]
    for r in rows[1:]:
        if r != prev + 1:
            runs.append((start, prev)); start = r
        prev = r
    runs.append((start, prev))
    blocks = []
    for r1, r2 in runs:
        cols = sorted({c for (r, c) in grid if r1 <= r <= r2})
        cstart, cprev = cols[0], cols[0]
        for c in cols[1:]:
            if c != cprev + 1:
                blocks.append((r1, r2, cstart, cprev)); cstart = c
            cprev = c
        blocks.append((r1, r2, cstart, cprev))
    return blocks


def analyze_block(sheet: str, grid, r1: int, r2: int, c1: int, c2: int, idx: int) -> Table:
    rows = list(range(r1, r2 + 1))
    cells_of = {r: [grid[(r, c)] for c in range(c1, c2 + 1) if (r, c) in grid] for r in rows}
    maxw = max(len(v) for v in cells_of.values())
    strong = weak = None
    for r in rows[:12]:
        cells = cells_of[r]
        n = len(cells)
        nstr = sum(1 for ci in cells if ci.is_str)
        nbold = sum(1 for ci in cells if ci.bold and ci.is_str)
        if n >= 2 and nstr / n >= 0.6 and n >= 0.5 * maxw:
            if nbold >= 2:
                strong = r; break
            if weak is None:
                weak = r
    header = strong if strong is not None else weak
    first_col_strings = all((cells_of[r][0].is_str if cells_of[r] else True) for r in rows)
    tid = f"{slug(sheet)}__b{idx}"
    # key-value block: two columns, labels on the left, no bold header row
    if maxw <= 2 and strong is None and first_col_strings and len(rows) >= 1 and any(len(cells_of[r]) == 2 for r in rows):
        t = Table(tid, tid, sheet, f"{sheet} inputs", None, "kv", c1, c2, r1, None, None, r2)
        for r in rows:
            cs = cells_of[r]
            if len(cs) == 1 and cs[0].is_str:
                t.title_rows.append(r); t.title = t.title or str(cs[0].value); continue
            if len(cs) == 2:
                k, v = cs
                t.pairs.append({"key": str(k.value), "value": iso(v.value), "address": v.addr, "formula": v.formula,
                                "comment": v.comment, "author": v.author, "input": v.formula is None and (v.is_num or v.is_date)})
        return t
    if header is None:
        t = Table(tid, tid, sheet, f"{sheet} notes", None, "notes", c1, c2, r1, None, None, r2)
        for r in rows:
            t.text.extend(str(ci.value) for ci in cells_of[r] if ci.is_str)
        return t
    t = Table(tid, tid, sheet, sheet, None, "table", c1, c2, r1, header, header + 1, r2)
    for r in rows:
        if r >= header:
            break
        cs = cells_of[r]
        if len(cs) <= 2 and all(ci.is_str for ci in cs):
            t.title_rows.append(r)
            if t.title is None and cs:
                t.title = str(cs[0].value)
        else:
            t.helper_rows.append(r)
    if header == r2:  # header only, no data
        t.first_data_row = None
    # derived rows (totals, subtotals, margin lines)
    for r in range(header + 1, r2 + 1):
        cs = cells_of[r]
        if not cs:
            continue
        label = next((str(ci.value) for ci in cs[:2] if ci.is_str), None)
        first_col_empty = (r, c1) not in grid
        if label and (TOTAL_RE.match(label) or (first_col_empty and DERIVED_RE.match(label))):
            t.derived_rows[r] = label
        elif first_col_empty and cs and all(ci.formula for ci in cs if ci.is_num or ci.formula) and label:
            t.derived_rows[r] = label
    return t


def type_columns(t: Table, grid) -> None:
    hdr = [grid.get((t.header_row, c)) for c in range(t.c1, t.c2 + 1)]
    data_rows = t.data_rows()
    seen: Counter = Counter()
    for i, c in enumerate(range(t.c1, t.c2 + 1)):
        name = str(hdr[i].value) if hdr[i] and hdr[i].value is not None else get_column_letter(c)
        cs = slug(name)
        seen[cs] += 1
        if seen[cs] > 1:
            cs = f"{cs}_{seen[cs]}"
        vals = [grid.get((r, c)) for r in data_rows]
        present = [ci for ci in vals if ci is not None and ci.value is not None]
        n = len(data_rows) or 1
        kinds = Counter()
        for ci in present:
            if ci.formula and ci.value is None:
                kinds["formula"] += 1
            elif isinstance(ci.value, bool):
                kinds["boolean"] += 1
            elif ci.is_num:
                kinds["integer" if float(ci.value).is_integer() and "." not in ci.fmt else "decimal"] += 1
            elif ci.is_date:
                kinds["date"] += 1
            else:
                kinds["text"] += 1
        base = kinds.most_common(1)[0][0] if kinds else "empty"
        if base == "formula":
            fs = [ci.formula for ci in present if ci.formula]
            base = "text" if any("TEXT(" in f or "&" in f for f in fs) else "decimal"
        fmts = Counter(ci.fmt for ci in present if not ci.is_str)
        fmt = fmts.most_common(1)[0][0] if fmts else "General"
        semantic = base
        if base in ("integer", "decimal"):
            if "%" in fmt:
                semantic = "percent"
            elif "$" in fmt or "#,##0" in fmt:
                semantic = "money"
        strvals = [str(ci.value).strip() for ci in present if ci.value is not None]
        distinct = len(set(strvals))
        distinct_ratio = distinct / len(strvals) if strvals else 0.0
        avg_len = sum(len(s) for s in strvals) / len(strvals) if strvals else 0
        is_key = bool(strvals) and len(strvals) >= 3 and distinct_ratio >= 0.98 and base in ("text", "integer") and avg_len <= 24 and semantic not in ("money", "percent")
        if base == "text" and strvals and (is_key or all(CODE_RE.match(s) for s in strvals[:50])) and avg_len <= 12:
            semantic = "identifier"
        if base == "text" and avg_len > 60:
            semantic = "narrative"; t.narrative = True
        formula_share = sum(1 for ci in present if ci.formula) / len(present) if present else 0.0
        t.columns.append({
            "id": f"{t.id}.{cs}", "slug": cs, "name": name, "letter": get_column_letter(c), "index": c,
            "type": base, "semantic": semantic, "number_format": fmt, "null_rate": round(1 - len(present) / n, 3),
            "distinct": distinct, "distinct_ratio": round(distinct_ratio, 3), "is_key": is_key,
            "formula_share": round(formula_share, 3), "samples": [iso(ci.value) for ci in present[:4]],
        })
    allcells = [grid.get((r, c)) for r in data_rows for c in range(t.c1, t.c2 + 1)]
    allcells = [ci for ci in allcells if ci is not None]
    t.formula_share = round(sum(1 for ci in allcells if ci.formula) / len(allcells), 3) if allcells else 0.0
    t.rows = [[iso(grid[(r, c)].value) if (r, c) in grid else None for c in range(t.c1, t.c2 + 1)] for r in data_rows]
    keys = [c for c in t.columns if c["is_key"]]
    if t.formula_share >= 0.5:
        t.role = "report"
    elif t.narrative:
        t.role = "narrative"
    elif re.search(r"assum|input|param|driver", t.sheet, re.I):
        t.role = "assumptions"
    elif keys and len(t.rows) <= 500:
        t.role = "dimension"
    else:
        t.role = "data"


def name_table(t: Table, sheet_tables: list[Table]) -> None:
    n_tables = sum(1 for x in sheet_tables if x.kind == "table")
    base = slug(t.sheet)
    if t.kind == "table":
        first = next((x for x in sheet_tables if x.kind == "table"), None)
        if first is not t:
            short = "_".join(slug(t.title).split("_")[:4]) if t.title else ""
            base = f"{slug(t.sheet)}__{short}" if short else f"{slug(t.sheet)}__r{t.header_row}"
        t.name = (t.title if (first is not t and t.title) else t.sheet)
    elif t.kind == "kv":
        base = f"{slug(t.sheet)}__inputs_r{t.top_row}"
    else:
        base = f"{slug(t.sheet)}__notes_r{t.top_row}"
    t.slug = base
    t.id = base
    for c in t.columns:
        c["id"] = f"{t.id}.{c['slug']}"


# ----------------------------------------------------------------------------- lineage

def parse_ref(value: str, cur_sheet: str, names: dict[str, str], depth: int = 0) -> dict | None:
    v = value.strip()
    ext = None
    if "!" in v:
        sheet, ref = v.rsplit("!", 1)
        sheet = sheet.strip("'")
        m = re.match(r"^\[([^\]]+)\](.*)$", sheet)
        if m:
            ext, sheet = m.group(1), m.group(2)
        if ":" in sheet:
            return None  # 3-D reference
    else:
        sheet, ref = cur_sheet, v
    ref = ref.replace("$", "").upper() if not NAME_RE.match(ref) or CELL_RE.match(ref.replace("$", "").upper()) else ref
    if "[" in ref or "#" in ref:
        return None
    if NAME_RE.match(ref) and not CELL_RE.match(ref) and not COLS_RE.match(ref):
        target = names.get(ref) or names.get(ref.upper())
        if not target or depth > 3:
            return {"kind": "name", "name": ref, "unresolved": True, "text": value}
        r = parse_ref(target, cur_sheet, names, depth + 1)
        if r:
            r["via_name"] = ref
        return r
    m = CELL_RE.match(ref)
    if m:
        c = column_index_from_string(m.group(1)); r = int(m.group(2))
        return {"kind": "range", "sheet": sheet, "ext": ext, "c1": c, "r1": r, "c2": c, "r2": r, "whole_col": False, "text": value}
    m = RANGE_RE.match(ref)
    if m:
        c1, r1, c2, r2 = column_index_from_string(m.group(1)), int(m.group(2)), column_index_from_string(m.group(3)), int(m.group(4))
        return {"kind": "range", "sheet": sheet, "ext": ext, "c1": min(c1, c2), "r1": min(r1, r2), "c2": max(c1, c2), "r2": max(r1, r2), "whole_col": False, "text": value}
    m = COLS_RE.match(ref)
    if m:
        return {"kind": "range", "sheet": sheet, "ext": ext, "c1": column_index_from_string(m.group(1)), "r1": 1,
                "c2": column_index_from_string(m.group(2)), "r2": None, "whole_col": True, "text": value}
    m = ROWS_RE.match(ref)
    if m:
        return {"kind": "range", "sheet": sheet, "ext": ext, "c1": 1, "r1": int(m.group(1)), "c2": None, "r2": int(m.group(2)), "whole_col": False, "whole_row": True, "text": value}
    return None


def top_level_calls(tok: Tokenizer) -> list[tuple[str, list[str]]]:
    """For each function call, its argument texts (nested calls kept verbatim). Returns [(FUNC, [arg, ...])]."""
    calls: list[tuple[str, list[str]]] = []
    stack: list[tuple[str, list[str], list[str]]] = []
    for t in tok.items:
        if t.type == "FUNC" and t.subtype == "OPEN":
            for s_ in stack:
                s_[2].append(t.value)
            stack.append((t.value.rstrip("(").upper(), [], []))
        elif t.type == "FUNC" and t.subtype == "CLOSE":
            if not stack:
                continue
            name, args, cur = stack.pop()
            args.append("".join(cur).strip())
            calls.append((name, args))
            for s_ in stack:
                s_[2].append(")")
        elif t.type == "SEP" and t.subtype == "ARG":
            if stack:
                name, args, cur = stack[-1]
                args.append("".join(cur).strip()); cur.clear()
                for s_ in stack[:-1]:
                    s_[2].append(",")
        else:
            for s_ in stack:
                s_[2].append(t.value)
    return calls


class Lineage:
    def __init__(self, grids, tables_by_sheet: dict[str, list[Table]], names: dict[str, str]):
        self.grids = grids
        self.tables_by_sheet = tables_by_sheet
        self.names = names
        self.nodes: dict[str, dict] = {}
        self.edges: dict[tuple[str, str, str], dict] = {}
        self.findings: list[dict] = []
        self.stale: dict[tuple[str, str, int], dict] = {}
        self.constants: dict[tuple[str, str], dict] = {}
        self.input_refs: Counter = Counter()

    # --- nodes
    def column_node(self, t: Table, col: int) -> str:
        c = next((x for x in t.columns if x["index"] == col), None)
        nid = c["id"] if c else f"{t.id}.{get_column_letter(col)}"
        if nid not in self.nodes:
            self.nodes[nid] = {"id": nid, "type": "COLUMN", "name": f"{t.name}.{c['name'] if c else get_column_letter(col)}",
                               "sheet": t.sheet, "table": t.id, "column": c["slug"] if c else get_column_letter(col)}
        return nid

    def row_node(self, t: Table, row: int) -> str:
        label = t.derived_rows.get(row, f"row {row}")
        nid = f"{t.id}#r{row}"
        if nid not in self.nodes:
            self.nodes[nid] = {"id": nid, "type": "METRIC_ROW", "name": f"{t.name}: {label}", "sheet": t.sheet, "table": t.id, "row": row, "label": label}
        return nid

    def cell_node(self, sheet: str, row: int, col: int) -> str:
        addr = f"{get_column_letter(col)}{row}"
        nid = f"{slug(sheet)}!{addr}"
        if nid in self.nodes:
            return nid
        grid = self.grids.get(sheet, {})
        ci = grid.get((row, col))
        label = None
        for t in self.tables_by_sheet.get(sheet, []):
            if t.kind == "kv":
                for p in t.pairs:
                    if p["address"] == addr:
                        label = p["key"]
            elif t.kind == "table" and row in t.helper_rows and t.c1 <= col <= t.c2:
                hdr = grid.get((t.header_row, col))
                label = f"{t.name} key for {hdr.value if hdr else get_column_letter(col)}"
            elif t.kind == "table" and t.first_data_row is not None and t.first_data_row <= row <= t.last_row and t.c1 <= col <= t.c2:
                hdr = grid.get((t.header_row, col)); key = grid.get((row, t.c1))
                label = f"{key.value if key and key.is_str else 'row ' + str(row)} / {hdr.value if hdr else get_column_letter(col)}"
        if label is None:
            for cc in range(col - 1, max(col - 4, 0), -1):
                left = grid.get((row, cc))
                if left is not None and left.is_str:
                    label = str(left.value); break
        value = iso(ci.value) if ci else None
        is_input = bool(ci and ci.formula is None and (ci.is_num or ci.is_date))
        self.nodes[nid] = {"id": nid, "type": "CELL", "name": f"{sheet}!{addr}" + (f" {label}" if label else "") + (f" = {value}" if value is not None and is_input else ""),
                           "sheet": sheet, "address": addr, "label": label, "value": value, "is_input": is_input,
                           "has_formula": bool(ci and ci.formula), "comment": ci.comment if ci else None}
        return nid

    def external_node(self, ext: str, sheet: str, text: str) -> str:
        nid = f"external:{slug(ext)}"
        if nid not in self.nodes:
            self.nodes[nid] = {"id": nid, "type": "EXTERNAL_WORKBOOK", "name": f"{ext} ({sheet})", "ref": text}
        return nid

    def nodes_for_ref(self, ref: dict, ctx_sheet: str, ctx_addr: str) -> list[str]:
        m = CELL_RE.match(ctx_addr)
        ctx_col, ctx_row = (column_index_from_string(m.group(1)), int(m.group(2))) if m else (0, 0)
        if ref["kind"] == "name":
            nid = f"name:{ref['name']}"
            self.nodes.setdefault(nid, {"id": nid, "type": "NAMED_RANGE", "name": ref["name"], "unresolved": True})
            return [nid]
        if ref.get("ext"):
            return [self.external_node(ref["ext"], ref["sheet"], ref["text"])]
        sheet = ref["sheet"]
        if sheet not in self.grids:
            match = next((s for s in self.grids if s.lower() == sheet.lower()), None)
            if match is None:
                return []
            sheet = match
        grid = self.grids[sheet]
        maxrow = max((r for r, _ in grid), default=1)
        maxcol = max((c for _, c in grid), default=1)
        r1, r2 = ref["r1"], ref["r2"] if ref["r2"] is not None else maxrow
        c1, c2 = ref["c1"], ref["c2"] if ref["c2"] is not None else maxcol
        single = (r1 == r2 and c1 == c2)
        out: list[str] = []
        for t in self.tables_by_sheet.get(sheet, []):
            if t.kind != "table" or t.first_data_row is None:
                continue
            if r2 < t.top_row or r1 > t.last_row or c2 < t.c1 or c1 > t.c2:
                continue
            cols = range(max(c1, t.c1), min(c2, t.c2) + 1)
            if single:
                same_table = (ctx_sheet == sheet and t.c1 <= ctx_col <= t.c2 and t.top_row <= ctx_row <= t.last_row)
                if t.first_data_row <= r1 <= t.last_row:
                    if r1 in t.derived_rows:
                        out.append(self.row_node(t, r1))
                    elif not same_table and (t.role == "assumptions" or (len(t.rows) <= 12 and c1 > t.c1)):
                        out.append(self.cell_node(sheet, r1, c1))
                    else:
                        out.append(self.column_node(t, c1))
                else:
                    out.append(self.cell_node(sheet, r1, c1))
                continue
            data_hit = any(r not in t.derived_rows for r in range(max(r1, t.first_data_row), min(r2, t.last_row) + 1))
            if data_hit:
                out.extend(self.column_node(t, c) for c in cols)
            for r in t.derived_rows:
                if r1 <= r <= r2:
                    out.append(self.row_node(t, r))
            same_table = (ctx_sheet == sheet and t.c1 <= ctx_col <= t.c2 and t.top_row <= ctx_row <= t.last_row)
            if (not ref["whole_col"] and not ref.get("whole_row") and not same_table and r2 < t.last_row
                    and r1 <= t.first_data_row + 1 and (r2 - r1) >= 10):
                key = (sheet, t.id, r2)
                f = self.stale.setdefault(key, {"kind": "stale_range", "severity": "high", "sheet": sheet, "table": t.id,
                                                "range_end": r2, "table_end": t.last_row, "rows_excluded": t.last_row - r2,
                                                "excluded_data_rows": sum(1 for r in t.data_rows() if r > r2), "formulas": 0, "examples": []})
                f["formulas"] += 1
                if len(f["examples"]) < 5 and f"{ctx_sheet}!{ctx_addr}" not in f["examples"]:
                    f["examples"].append(f"{ctx_sheet}!{ctx_addr}")
        if not out:
            if single:
                out.append(self.cell_node(sheet, r1, c1))
            else:
                for r in range(r1, min(r2, r1 + 50) + 1):
                    for c in range(c1, min(c2, c1 + 20) + 1):
                        if (r, c) in grid:
                            out.append(self.cell_node(sheet, r, c))
        return out

    def target_for(self, sheet: str, ci: CellInfo) -> str:
        for t in self.tables_by_sheet.get(sheet, []):
            if t.kind == "table" and t.first_data_row is not None and t.c1 <= ci.col <= t.c2 and t.first_data_row <= ci.row <= t.last_row:
                return self.row_node(t, ci.row) if ci.row in t.derived_rows else self.column_node(t, ci.col)
        return self.cell_node(sheet, ci.row, ci.col)

    def add_edge(self, src: str, dst: str, pred: str, formula: str, where: str, via: str | None = None) -> None:
        if src == dst:
            return
        e = self.edges.setdefault((src, dst, pred), {"src": src, "dst": dst, "predicate": pred, "count": 0, "context": formula[:200], "cells": []})
        e["count"] += 1
        if via:
            e["via_name"] = via
        if len(e["cells"]) < 5:
            e["cells"].append(where)

    def run(self) -> None:
        for sheet, grid in self.grids.items():
            for ci in list(grid.values()):
                if not ci.formula:
                    continue
                try:
                    tok = Tokenizer(ci.formula)
                except Exception:
                    continue
                funcs = [t.value.rstrip("(").upper() for t in tok.items if t.type == "FUNC" and t.subtype == "OPEN"]
                refs = [t.value for t in tok.items if t.type == "OPERAND" and t.subtype == "RANGE"]
                nums = [t.value for t in tok.items if t.type == "OPERAND" and t.subtype == "NUMBER"]
                ops = [t for t in tok.items if t.type == "OPERATOR-INFIX"]
                pred = funcs[0] if funcs else ("ARITH" if ops else "REF")
                where = f"{sheet}!{ci.addr}"
                dst = self.target_for(sheet, ci)
                handled: set[str] = set()
                for fname, args in top_level_calls(tok):
                    if fname in ("VLOOKUP", "HLOOKUP") and len(args) >= 3:
                        key_ref = parse_ref(args[0], sheet, self.names) if args[0] else None
                        arr_ref = parse_ref(args[1], sheet, self.names) if args[1] else None
                        try:
                            idx = int(float(args[2])) - 1
                        except ValueError:
                            idx = None
                        if key_ref and arr_ref and arr_ref["kind"] == "range":
                            handled.update({args[0], args[1]})
                            keys = self.nodes_for_ref(key_ref, sheet, ci.addr)
                            first = dict(arr_ref, c2=arr_ref["c1"]) if fname == "VLOOKUP" else dict(arr_ref, r2=arr_ref["r1"])
                            firsts = self.nodes_for_ref(first, sheet, ci.addr)
                            for k in keys:
                                for f in firsts:
                                    self.add_edge(k, f, "LOOKS_UP", ci.formula, where)
                                self.add_edge(k, dst, f"{fname}_KEY", ci.formula, where)
                            if idx is not None and arr_ref["c2"] is not None:
                                ret = dict(arr_ref, c1=arr_ref["c1"] + idx, c2=arr_ref["c1"] + idx) if fname == "VLOOKUP" else dict(arr_ref, r1=arr_ref["r1"] + idx, r2=arr_ref["r1"] + idx)
                                for src in self.nodes_for_ref(ret, sheet, ci.addr):
                                    self.add_edge(src, dst, fname, ci.formula, where)
                    elif fname == "XLOOKUP" and len(args) >= 3:
                        key_ref, arr_ref, ret_ref = (parse_ref(a, sheet, self.names) if a else None for a in args[:3])
                        if key_ref and arr_ref and ret_ref:
                            handled.update(args[:3])
                            keys = self.nodes_for_ref(key_ref, sheet, ci.addr)
                            for k in keys:
                                for f in self.nodes_for_ref(arr_ref, sheet, ci.addr):
                                    self.add_edge(k, f, "LOOKS_UP", ci.formula, where)
                                self.add_edge(k, dst, "XLOOKUP_KEY", ci.formula, where)
                            for src in self.nodes_for_ref(ret_ref, sheet, ci.addr):
                                self.add_edge(src, dst, "XLOOKUP", ci.formula, where)
                for rv in refs:
                    if rv in handled:
                        continue
                    ref = parse_ref(rv, sheet, self.names)
                    if not ref:
                        continue
                    for src in self.nodes_for_ref(ref, sheet, ci.addr):
                        self.add_edge(src, dst, pred, ci.formula, where, via=ref.get("via_name"))
                        n = self.nodes.get(src)
                        if n and n.get("type") == "CELL" and n.get("is_input"):
                            self.input_refs[src] += 1
                # findings on the formula itself
                for v in nums:
                    try:
                        x = float(v)
                    except ValueError:
                        continue
                    if x in BENIGN_CONSTANTS:
                        continue
                    if "." in v or abs(x) >= 1000:
                        f = self.constants.setdefault((sheet, v), {"kind": "constant_in_formula", "severity": "medium", "sheet": sheet, "location": ci.addr,
                                                                    "detail": f"literal {v} typed inside formulas, e.g. {ci.formula[:80]}", "count": 0, "examples": []})
                        f["count"] += 1
                        if len(f["examples"]) < 5:
                            f["examples"].append(ci.addr)
                up = ci.formula.upper()
                if "VLOOKUP(" in up or "HLOOKUP(" in up:
                    args = up.split("VLOOKUP(", 1)[-1] if "VLOOKUP(" in up else up.split("HLOOKUP(", 1)[-1]
                    if not re.search(r",\s*(FALSE|0)\s*\)", args):
                        self.findings.append({"kind": "approximate_lookup", "severity": "medium", "sheet": sheet, "location": ci.addr,
                                              "detail": "VLOOKUP without exact-match flag (FALSE/0)"})
                if "[" in ci.formula:
                    self.findings.append({"kind": "external_link", "severity": "high", "sheet": sheet, "location": ci.addr, "detail": ci.formula[:120]})
        self.findings.extend(self.stale.values())
        self.findings.extend(self.constants.values())


# ----------------------------------------------------------------------------- relationships

def infer_relationships(tables: list[Table]) -> list[dict]:
    rels = []
    keyed = []
    for t in tables:
        if t.kind != "table" or t.role == "report":
            continue
        for c in t.columns:
            if c["is_key"] and c["type"] in ("text", "integer") and len(t.rows) >= 2 and c["formula_share"] < 0.5:
                vals = {str(r[c["index"] - t.c1]).strip() for r in t.rows if r[c["index"] - t.c1] is not None}
                keyed.append((t, c, vals))
    for t in tables:
        if t.kind != "table":
            continue
        for c in t.columns:
            if c["type"] not in ("text", "integer") or c["semantic"] == "narrative" or c["formula_share"] >= 0.5:
                continue
            vals = {str(r[c["index"] - t.c1]).strip() for r in t.rows if r[c["index"] - t.c1] is not None}
            if len(vals) < 2:
                continue
            best = None
            for pt, pc, pvals in keyed:
                if pt.id == t.id or (c["is_key"] and len(vals) >= len(pvals)):
                    continue
                cover = len(vals & pvals) / len(vals)
                if cover >= 0.9 and len(pvals) >= 2:
                    score = (cover, pc["semantic"] == "identifier", pt.role == "dimension", -len(pvals))
                    if best is None or score > best[0]:
                        best = (score, pt, pc, cover)
            if best:
                _, pt, pc, cover = best
                rels.append({"child_table": t.id, "child_col": c["id"], "parent_table": pt.id, "parent_col": pc["id"],
                             "kind": "fk_inferred", "coverage": round(cover, 3), "parent_role": pt.role})
    return rels


def cross_workbook_relationships(captures: list[dict]) -> list[dict]:
    """Foreign keys between tables of different workbooks (same rule as infer_relationships)."""
    rels = []
    keyed = []
    for cap in captures:
        for t in cap["tables"]:
            if t["kind"] != "table" or t["role"] == "report":
                continue
            for c in t["columns"]:
                if c["is_key"] and c["type"] in ("text", "integer") and len(t["rows"]) >= 2 and c["formula_share"] < 0.5:
                    i = c["index"] - t["c1"]
                    keyed.append((cap["workbook"]["slug"], t, c, {str(r[i]).strip() for r in t["rows"] if r[i] is not None}))
    for cap in captures:
        for t in cap["tables"]:
            if t["kind"] != "table":
                continue
            for c in t["columns"]:
                if c["type"] not in ("text", "integer") or c["semantic"] == "narrative" or c["formula_share"] >= 0.5:
                    continue
                i = c["index"] - t["c1"]
                vals = {str(r[i]).strip() for r in t["rows"] if r[i] is not None}
                if len(vals) < 2:
                    continue
                for pwb, pt, pc, pvals in keyed:
                    if pwb == cap["workbook"]["slug"] or (c["is_key"] and len(vals) >= len(pvals)):
                        continue
                    cover = len(vals & pvals) / len(vals)
                    if cover >= 0.9:
                        rels.append({"child_workbook": cap["workbook"]["slug"], "child_table": t["id"], "child_col": c["id"],
                                     "parent_workbook": pwb, "parent_table": pt["id"], "parent_col": pc["id"], "kind": "fk_cross_workbook", "coverage": round(cover, 3)})
    return rels


# ----------------------------------------------------------------------------- notes, scorecard, DDL

def collect_notes(grids, tables_by_sheet) -> list[dict]:
    notes = []
    for sheet, grid in grids.items():
        for ci in grid.values():
            if ci.comment:
                notes.append({"sheet": sheet, "address": ci.addr, "kind": "comment", "text": ci.comment.strip(), "author": ci.author,
                              "provenance": bool(PASTED_RE.search(ci.comment))})
        for t in tables_by_sheet.get(sheet, []):
            if t.kind == "notes":
                for line in t.text:
                    notes.append({"sheet": sheet, "address": t.range, "kind": "text", "text": line, "author": None, "provenance": bool(PASTED_RE.search(line))})
            elif t.kind == "kv":
                for p in t.pairs:
                    notes.append({"sheet": sheet, "address": p["address"], "kind": "fact", "text": f"{p['key']}: {p['value']}", "author": None, "provenance": False})
            for r in t.title_rows:
                for c in range(t.c1, t.c2 + 1):
                    ci = grid.get((r, c))
                    if ci and ci.is_str:
                        notes.append({"sheet": sheet, "address": ci.addr, "kind": "title", "text": str(ci.value), "author": None, "provenance": bool(PASTED_RE.search(str(ci.value)))})
    return notes


SQL_TYPES = {"integer": "BIGINT", "decimal": "DECIMAL(18,2)", "money": "DECIMAL(18,2)", "percent": "DECIMAL(9,6)", "date": "DATE",
             "boolean": "BOOLEAN", "text": "VARCHAR", "identifier": "VARCHAR", "narrative": "VARCHAR", "empty": "VARCHAR"}


def ddl_for(wb_slug: str, t: Table, rels: list[dict]) -> str:
    cols = []
    for i, c in enumerate(t.columns):
        typ = SQL_TYPES.get(c["semantic"], SQL_TYPES.get(c["type"], "VARCHAR"))
        last = i == len(t.columns) - 1
        cols.append(f"  {c['slug']} {typ}" + ("" if last else ",") + ("  -- candidate key" if c["is_key"] else ""))
    fks = [f"-- FK {r['child_col'].split('.')[-1]} -> {r['parent_table']}.{r['parent_col'].split('.')[-1]} (coverage {r['coverage']})" for r in rels if r["child_table"] == t.id]
    head = f"-- {t.sheet} / {t.name}: {len(t.rows)} rows, role={t.role}, formula share={t.formula_share}\n"
    if t.role == "report":
        head += "-- report table: re-derive as a view or metric; the DDL below stores its current values\n"
    head += ("\n".join(fks) + "\n") if fks else ""
    return head + f"CREATE TABLE {wb_slug}__{t.slug} (\n" + "\n".join(cols) + "\n);\n"


def scorecard(grids, sheets_meta, tables: list[Table], lineage: Lineage, notes: list[dict], names: dict) -> dict:
    total = sum(len(g) for g in grids.values())
    formulas = sum(1 for g in grids.values() for ci in g.values() if ci.formula)
    data_cells = sum(1 for t in tables if t.kind == "table" and t.role in ("data", "dimension", "assumptions")
                     for r in t.data_rows() for c in range(t.c1, t.c2 + 1) if (r, c) in grids[t.sheet] and not grids[t.sheet][(r, c)].formula)
    inputs = [n for n in lineage.nodes.values() if n["type"] == "CELL" and n.get("is_input")]
    labels = total - formulas - data_cells
    findings = list(lineage.findings)
    for n in notes:
        if n["provenance"]:
            findings.append({"kind": "pasted_value", "severity": "medium", "sheet": n["sheet"], "location": n["address"], "detail": n["text"][:160]})
    for s in sheets_meta:
        if s["hidden_rows"]:
            findings.append({"kind": "hidden_rows", "severity": "low", "sheet": s["name"], "location": ",".join(map(str, s["hidden_rows"][:10])), "detail": f"{len(s['hidden_rows'])} hidden row(s); formulas may depend on them"})
        if s["hidden_cols"]:
            findings.append({"kind": "hidden_cols", "severity": "low", "sheet": s["name"], "location": ",".join(s["hidden_cols"][:10]), "detail": f"{len(s['hidden_cols'])} hidden column(s)"})
    for n in inputs:
        if lineage.input_refs.get(n["id"]) and " key for " not in (n.get("label") or ""):
            findings.append({"kind": "manual_input", "severity": "info", "sheet": n["sheet"], "location": n["address"],
                             "detail": f"{n.get('label') or n['address']} = {n.get('value')}; referenced by {lineage.input_refs[n['id']]} formula(s)" + (f"; note: {n['comment'][:120]}" if n.get("comment") else "")})
    sev = Counter(f["severity"] for f in findings)
    roles = Counter(t.role for t in tables if t.kind == "table")
    return {
        "cells": total, "formula_cells": formulas, "data_cells": max(data_cells, 0), "label_or_presentation_cells": max(labels, 0),
        "formula_share": round(formulas / total, 3) if total else 0, "named_ranges": len(names),
        "manual_inputs": len([n for n in inputs if lineage.input_refs.get(n["id"])]),
        "tables_by_role": dict(roles), "findings_by_severity": dict(sev), "findings": findings,
        "migration_plan": {
            "load_as_tables": [t.id for t in tables if t.kind == "table" and t.role in ("data", "dimension")],
            "rebuild_as_views_or_metrics": [t.id for t in tables if t.kind == "table" and t.role == "report"],
            "keep_as_assumptions": [t.id for t in tables if t.kind == "kv"],
            "ingest_as_knowledge": [t.id for t in tables if t.kind == "notes" or (t.kind == "table" and t.role == "narrative")],
        },
    }


# ----------------------------------------------------------------------------- entry point

def capture(path: str | Path, workbook_slug: str | None = None) -> dict:
    path = Path(path)
    grids, sheets_meta, names = load_grid(path)
    wb_slug = workbook_slug or slug(path.stem)
    tables_by_sheet: dict[str, list[Table]] = {}
    for sheet, grid in grids.items():
        ts = []
        for i, (r1, r2, c1, c2) in enumerate(find_blocks(grid), start=1):
            t = analyze_block(sheet, grid, r1, r2, c1, c2, i)
            if t.kind == "table":
                type_columns(t, grid)
            ts.append(t)
        for i, t in enumerate(ts):
            if t.kind == "table" and t.title is None and i > 0 and ts[i - 1].kind == "notes" and len(ts[i - 1].text) == 1 and ts[i - 1].last_row >= t.top_row - 2:
                t.title = ts[i - 1].text[0]
        for t in ts:
            name_table(t, ts)
        tables_by_sheet[sheet] = ts
    tables = [t for ts in tables_by_sheet.values() for t in ts]
    lin = Lineage(grids, tables_by_sheet, names)
    lin.run()
    # explicit lookup relationships from VLOOKUP/XLOOKUP edges
    rels = infer_relationships(tables)
    seen_rel = {(r["child_col"], r["parent_col"]) for r in rels}
    for (src, dst, pred), e in lin.edges.items():
        if pred == "LOOKS_UP":
            s, d = lin.nodes.get(src, {}), lin.nodes.get(dst, {})
            if s.get("type") == "COLUMN" and d.get("type") == "COLUMN" and s.get("table") != d.get("table") and (src, dst) not in seen_rel:
                rels.append({"child_table": s["table"], "child_col": src, "parent_table": d["table"], "parent_col": dst, "kind": "lookup_formula", "coverage": 1.0, "parent_role": ""})
                seen_rel.add((src, dst))
    notes = collect_notes(grids, tables_by_sheet)
    sc = scorecard(grids, sheets_meta, tables, lin, notes, names)
    ddl = "\n".join(ddl_for(wb_slug, t, rels) for t in tables if t.kind == "table" and t.columns)
    # workbook / sheet / table containment edges for the graph
    nodes = dict(lin.nodes)
    nodes[f"wb:{wb_slug}"] = {"id": f"wb:{wb_slug}", "type": "WORKBOOK", "name": path.name}
    edges = list(lin.edges.values())
    for sheet in grids:
        sid = f"sheet:{slug(sheet)}"
        nodes[sid] = {"id": sid, "type": "SHEET", "name": sheet}
        edges.append({"src": f"wb:{wb_slug}", "dst": sid, "predicate": "CONTAINS", "count": 1, "context": "", "cells": []})
        for t in tables_by_sheet[sheet]:
            if t.kind != "table":
                continue
            tid = f"table:{t.id}"
            nodes[tid] = {"id": tid, "type": "TABLE", "name": t.name, "sheet": sheet, "table": t.id, "role": t.role}
            edges.append({"src": sid, "dst": tid, "predicate": "CONTAINS", "count": 1, "context": "", "cells": []})
            for c in t.columns:
                nodes.setdefault(c["id"], {"id": c["id"], "type": "COLUMN", "name": f"{t.name}.{c['name']}", "sheet": sheet, "table": t.id, "column": c["slug"]})
                edges.append({"src": tid, "dst": c["id"], "predicate": "HAS_COLUMN", "count": 1, "context": "", "cells": []})
    for r in rels:
        if r["kind"].startswith("fk"):
            edges.append({"src": r["child_col"], "dst": r["parent_col"], "predicate": "REFERENCES", "count": 1, "context": f"{r['kind']} coverage {r['coverage']}", "cells": []})
    out = {
        "workbook": {"slug": wb_slug, "filename": path.name, "path": str(path), "sheets": sheets_meta, "named_ranges": names,
                     "captured_at": dt.datetime.now().isoformat(timespec="seconds")},
        "tables": [table_dict(t) for t in tables],
        "nodes": nodes, "edges": edges, "relationships": rels, "notes": notes, "findings": sc["findings"], "scorecard": sc, "ddl": ddl,
    }
    return out


def table_dict(t: Table) -> dict:
    return {"id": t.id, "slug": t.slug, "sheet": t.sheet, "name": t.name, "title": t.title, "kind": t.kind, "role": t.role, "range": t.range,
            "c1": t.c1, "c2": t.c2, "top_row": t.top_row, "header_row": t.header_row, "first_data_row": t.first_data_row, "last_row": t.last_row,
            "helper_rows": t.helper_rows, "title_rows": t.title_rows, "derived_rows": {str(k): v for k, v in t.derived_rows.items()},
            "columns": t.columns, "n_rows": len(t.rows), "rows": t.rows, "pairs": t.pairs, "text": t.text,
            "formula_share": t.formula_share, "narrative": t.narrative}


def summary(cap: dict) -> str:
    lines = [f"{cap['workbook']['filename']}: {len(cap['tables'])} blocks, {len(cap['edges'])} edges, {len(cap['relationships'])} relationships, {len(cap['findings'])} findings"]
    for t in cap["tables"]:
        extra = f" rows={t['n_rows']} cols={len(t['columns'])} role={t['role']} fshare={t['formula_share']}" if t["kind"] == "table" else (f" pairs={len(t['pairs'])}" if t["kind"] == "kv" else f" lines={len(t['text'])}")
        derived = f" derived={list(t['derived_rows'].values())[:4]}" if t["derived_rows"] else ""
        lines.append(f"  [{t['kind']:5}] {t['sheet']!r:16} {t['range']:10} hdr={t['header_row']} {t['id']}{extra}{derived}")
    for r in cap["relationships"]:
        lines.append(f"  rel {r['kind']}: {r['child_col']} -> {r['parent_col']} ({r['coverage']})")
    byk = Counter(f["kind"] for f in cap["findings"])
    lines.append("  findings: " + ", ".join(f"{k}={v}" for k, v in byk.items()))
    return "\n".join(lines)


if __name__ == "__main__":
    import sys
    for p in sys.argv[1:]:
        cap = capture(p)
        print(summary(cap))
        out = Path("out") / cap["workbook"]["slug"]
        out.mkdir(parents=True, exist_ok=True)
        (out / "capture.json").write_text(json.dumps(cap, indent=1, default=str))
        (out / "ddl.sql").write_text(cap["ddl"])
